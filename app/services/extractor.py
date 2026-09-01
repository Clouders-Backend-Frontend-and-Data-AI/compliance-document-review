import io
import os
from typing import Optional
import pypdf
import docx
import openpyxl

class DocumentExtractor:
    """Multi-format document text extraction engine (PDF, DOCX, XLSX, TXT)"""

    @staticmethod
    def extract_text_from_file(file_path: str, mime_type: Optional[str] = None) -> str:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf" or mime_type == "application/pdf":
            return DocumentExtractor._extract_from_pdf(file_path)
        elif ext in [".docx", ".doc"] or mime_type in [
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword"
        ]:
            return DocumentExtractor._extract_from_docx(file_path)
        elif ext in [".xlsx", ".xls"] or mime_type in [
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel"
        ]:
            return DocumentExtractor._extract_from_xlsx(file_path)
        elif ext in [".txt", ".md", ".csv"]:
            return DocumentExtractor._extract_from_plain_text(file_path)
        else:
            # Attempt plain text read as fallback
            try:
                return DocumentExtractor._extract_from_plain_text(file_path)
            except Exception:
                raise ValueError(f"Unsupported file format: {ext}")

    @staticmethod
    def extract_text_from_bytes(content_bytes: bytes, filename: str, mime_type: Optional[str] = None) -> str:
        ext = os.path.splitext(filename)[1].lower()
        bio = io.BytesIO(content_bytes)

        if ext == ".pdf" or mime_type == "application/pdf":
            reader = pypdf.PdfReader(bio)
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n\n".join(pages).strip()
        elif ext in [".docx", ".doc"]:
            doc = docx.Document(bio)
            text_chunks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        text_chunks.append(row_text)
            return "\n\n".join(text_chunks).strip()
        elif ext in [".xlsx", ".xls"]:
            wb = openpyxl.load_workbook(bio, data_only=True)
            text_chunks = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                text_chunks.append(f"--- Sheet: {sheet} ---")
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if row_vals:
                        text_chunks.append(" | ".join(row_vals))
            return "\n".join(text_chunks).strip()
        else:
            return content_bytes.decode("utf-8", errors="ignore").strip()

    @staticmethod
    def _extract_from_pdf(file_path: str) -> str:
        reader = pypdf.PdfReader(file_path)
        pages = []
        for idx, page in enumerate(reader.pages):
            txt = page.extract_text()
            if txt:
                pages.append(f"[Page {idx+1}]\n{txt.strip()}")
        return "\n\n".join(pages).strip()

    @staticmethod
    def _extract_from_docx(file_path: str) -> str:
        doc = docx.Document(file_path)
        text_chunks = []
        for p in doc.paragraphs:
            if p.text.strip():
                text_chunks.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    text_chunks.append(row_text)
        return "\n\n".join(text_chunks).strip()

    @staticmethod
    def _extract_from_xlsx(file_path: str) -> str:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        text_chunks = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            text_chunks.append(f"--- Sheet: {sheet} ---")
            for row in ws.iter_rows(values_only=True):
                row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if row_vals:
                    text_chunks.append(" | ".join(row_vals))
        return "\n".join(text_chunks).strip()

    @staticmethod
    def _extract_from_plain_text(file_path: str) -> str:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
