"""Generate sample PDF / DOCX / XLSX fixtures for local testing."""
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "sample_test_files"


def _write_simple_pdf(path: Path, lines) -> None:
    text = " ".join(lines)[:800].replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET"
    objects = [
        "1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj",
        "2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj",
        (
            "3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>endobj"
        ),
        f"4 0 obj<< /Length {len(content)} >>stream\n{content}\nendstream endobj",
        "5 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj",
    ]
    pdf = "%PDF-1.4\n"
    offsets = [0]
    for obj in objects:
        offsets.append(len(pdf))
        pdf += obj + "\n"
    xref_pos = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n"
    pdf += "0000000000 65535 f \n"
    for off in offsets[1:]:
        pdf += f"{off:010d} 00000 n \n"
    pdf += (
        f"trailer<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    )
    path.write_bytes(pdf.encode("latin-1", errors="replace"))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    compliant = [
        "Prepared for: Alex Client",
        "Past performance is not indicative of future results.",
        "Investing in securities involves risk of loss.",
        "NOT FDIC INSURED MAY LOSE VALUE NOT BANK GUARANTEED.",
    ]
    noncompliant = [
        "Dear Jordan Smith,",
        "Alpha Crypto Growth Sleeve promo.",
        "Our AI strategy delivers a Guaranteed 18 percent annual return with zero risk!",
        "Contact jordan.smith@example.com or call 555-234-5678.",
        "Account ACCT-982341 balance $1,500,000.",
    ]
    _write_simple_pdf(OUT / "sample_compliant_brochure.pdf", compliant)
    _write_simple_pdf(OUT / "sample_noncompliant_promo.pdf", noncompliant)

    doc = Document()
    for p in [
        "Executive Summary",
        "Prepared for: Robert Henderson",
        "Casey Investor proposal overview.",
    ] + compliant:
        doc.add_paragraph(p)
    doc.save(OUT / "sample_proposal.docx")
    doc.save(OUT / "sample_compliant_proposal.docx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Performance"
    for r_idx, row in enumerate(
        [
            ["Item", "Detail"],
            ["Sleeve", "Core Fixed Income"],
            ["Client", "Riley Example"],
            ["Email", "riley@example.com"],
            ["Fee", "0.85% AUM"],
        ],
        start=1,
    ):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    wb.save(OUT / "sample_fee_schedule.xlsx")
    wb.save(OUT / "sample_performance_sheet.xlsx")
    print(f"Wrote sample files to {OUT}")


if __name__ == "__main__":
    main()
