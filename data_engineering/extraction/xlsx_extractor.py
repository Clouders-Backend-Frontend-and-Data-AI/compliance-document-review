import io
import structlog
import openpyxl
from .base import BaseExtractor, ExtractionResult

logger = structlog.get_logger(__name__)


class XLSXExtractor(BaseExtractor):
    """
    Extracts text from XLSX files, sheet by sheet, row by row.
    Captures financial figures and associated names — compliance-relevant content
    often lives in spreadsheet form.
    """

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        warnings: list[str] = []
        sheets_text: list[str] = []

        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except Exception as e:
            raise ValueError(f"Could not open XLSX '{filename}': {e}") from e

        metadata = {"sheet_names": wb.sheetnames}

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_text: list[str] = []

            for row in ws.iter_rows(values_only=True):
                cell_values = [str(cell) if cell is not None else "" for cell in row]
                row_text = " | ".join(cell_values).strip(" |")
                if row_text:
                    rows_text.append(row_text)

            if rows_text:
                sheet_block = f"[SHEET: {sheet_name}]\n" + "\n".join(rows_text)
                sheets_text.append(sheet_block)
            else:
                warnings.append(f"Sheet '{sheet_name}' contains no readable data.")

        wb.close()
        raw_text = "\n\n".join(sheets_text)

        if not raw_text.strip():
            warnings.append("All sheets are empty or contain no readable data.")

        logger.info(
            "xlsx_extracted",
            filename=filename,
            sheet_count=len(wb.sheetnames),
            char_count=len(raw_text),
        )

        return ExtractionResult(
            raw_text=raw_text,
            page_count=len(wb.sheetnames),
            metadata=metadata,
            warnings=warnings,
        )